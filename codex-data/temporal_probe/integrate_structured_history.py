"""整合受控规模的 SIGNOR、TRRUST 与 miRTarBase 历史调控边。

数据源职责：
* SIGNOR Oct2022：蛋白/基因间有方向的因果激活、抑制或一般调控；
* TRRUST v2 human：转录因子对靶基因的激活、抑制或未知方向调控；
* miRTarBase v9 strong：强实验支持的人类 miRNA -> 靶基因关系。

共同处理顺序：物种过滤 -> 实体规范化 -> PMID 年份过滤 -> 重复证据聚合
-> 每个源实体/数据库最多 50 条 -> 边表 -> 逐数据库节点特征。

注意：数据库 release 本身均不晚于 2022。PMID 年份已知且 >2022 的记录
必须剔除；年份缺失的记录仍可凭历史快照版本进入，但在边表中保留空 year，
便于更严格实验再次排除。
"""
from __future__ import annotations

import csv
import json
import re
import sqlite3
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "codex-data" / "temporal_probe" / "structured_sources"
GRAPH = ROOT / "codex-output" / "temporal_probe" / "normalized_graph"
HGNC = ROOT / "codex-data" / "temporal_probe" / "cache" / "hgnc_complete_set_current.txt"
YEAR_DB = ROOT / "codex-data" / "temporal_probe" / "cache" / "pubmed_years.sqlite"
# 防止少数枢纽调控因子垄断特征；后续规模消融可测试 20/50/100。
MAX_PER_SOURCE_DB = 50


def read_csv(path, delimiter=","):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def write_csv(path, rows, fields):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def split_multi(value):
    return [x.strip() for x in (value or "").split("|") if x.strip()]


def load_maps():
    """加载候选实体、HGNC 正式符号/别名、UniProt 及人工规范化决策。"""
    nodes = read_csv(GRAPH / "nodes.csv")
    candidates = {r["name"] for r in nodes if r["node_type"] not in {"Pathway"}}
    exact, aliases, uniprot = {}, defaultdict(set), {}
    with HGNC.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle, delimiter="\t"):
            if row.get("status") != "Approved":
                continue
            symbol = row["symbol"]
            exact[symbol.upper()] = symbol
            for value in split_multi(row.get("prev_symbol", "")) + split_multi(row.get("alias_symbol", "")):
                aliases[value.upper()].add(symbol)
            for value in split_multi(row.get("uniprot_ids", "")):
                uniprot[value] = symbol
    decisions = read_csv(GRAPH / "normalization_decisions.csv")
    for row in decisions:
        if row["normalization_status"] == "accepted":
            exact[row["original_symbol"].upper()] = row["canonical_symbol"]
    return candidates, exact, aliases, uniprot


def resolve_symbol(value, candidates, exact, aliases):
    """将名称解析为唯一 HGNC symbol，并要求其落在受控候选宇宙内。"""
    key = (value or "").strip().upper()
    symbol = exact.get(key)
    if not symbol and len(aliases.get(key, set())) == 1:
        symbol = next(iter(aliases[key]))
    return symbol if symbol in candidates else None


def pmid_years():
    """读取本地 PMID->year 缓存，避免在处理阶段反复调用网络 API。"""
    db = sqlite3.connect(YEAR_DB)
    return dict(db.execute("SELECT pmid,year FROM years"))


def normalize_effect(value):
    """把 SIGNOR 原始效应词归一为 ACTIVATES/INHIBITS/REGULATES。"""
    text = (value or "").lower()
    if "up-regulates" in text or "activation" in text:
        return "ACTIVATES"
    if "down-regulates" in text or "inhibition" in text:
        return "INHIBITS"
    return "REGULATES"


def signor_edges(candidates, exact, aliases, uniprot, years):
    """抽取人类 protein->protein SIGNOR 边并执行 PMID 截止线过滤。"""
    raw = read_csv(RAW / "SIGNOR_Oct2022.tsv", delimiter="\t")
    result = []
    for row in raw:
        if row.get("TAX_ID") not in {"9606", "-1"}:
            continue
        if row.get("TYPEA") != "protein" or row.get("TYPEB") != "protein":
            continue
        left = uniprot.get(row.get("IDA", "")) or resolve_symbol(row.get("ENTITYA"), candidates, exact, aliases)
        right = uniprot.get(row.get("IDB", "")) or resolve_symbol(row.get("ENTITYB"), candidates, exact, aliases)
        if left not in candidates or right not in candidates or left == right:
            continue
        pmid = re.sub(r"\D", "", row.get("PMID", ""))
        year = years.get(pmid)
        if year and year > 2022:
            continue
        result.append({
            "source": left, "target": right, "relation": normalize_effect(row.get("EFFECT")),
            "database": "SIGNOR", "pmid": pmid, "year": year or "",
            "evidence_level": "curated_causal_direct" if row.get("DIRECT") == "t" else "curated_causal",
            "mechanism": row.get("MECHANISM", ""), "direct": row.get("DIRECT", ""),
            "release": "Oct2022",
        })
    return result, len(raw)


def trrust_edges(candidates, exact, aliases, years):
    """抽取 TRRUST v2 人类 TF->target；每个 PMID 先作为独立证据记录。"""
    result, raw_count = [], 0
    with (RAW / "TRRUST_v2_human.tsv").open(encoding="utf-8-sig", newline="") as handle:
        for parts in csv.reader(handle, delimiter="\t"):
            if len(parts) < 4:
                continue
            raw_count += 1
            left = resolve_symbol(parts[0], candidates, exact, aliases)
            right = resolve_symbol(parts[1], candidates, exact, aliases)
            if not left or not right or left == right:
                continue
            relation = {
                "Activation": "ACTIVATES", "Repression": "INHIBITS",
                "Unknown": "REGULATES",
            }.get(parts[2], "REGULATES")
            for pmid in re.findall(r"\d+", parts[3]):
                year = years.get(pmid)
                if year and year > 2022:
                    continue
                result.append({
                    "source": left, "target": right, "relation": relation,
                    "database": "TRRUST", "pmid": pmid, "year": year or "",
                    "evidence_level": "manually_curated_transcriptional",
                    "mechanism": "transcriptional regulation", "direct": "",
                    "release": "v2_2018",
                })
    return result, raw_count


def mirna_symbol(name, candidates):
    """将 hsa-miR-* 名称映射至候选集中的 MIR 基因符号；歧义时拒绝。"""
    match = re.search(r"(?:HSA-)?MIR-?(\d+[A-Z]?)", (name or "").upper())
    if not match:
        return None
    root = "MIR" + match.group(1)
    matches = [g for g in candidates if g == root or g.startswith(root + "-")]
    return matches[0] if len(matches) == 1 else None


def mirtarbase_edges(candidates, exact, aliases, years):
    """抽取人类-人类且强实验支持的 miRNA->target Gene 关系。"""
    workbook = load_workbook(
        RAW / "miRTarBase_v9_strong_WR.xlsx", read_only=True, data_only=True
    )
    sheet = workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    header = [str(x) for x in next(iterator)]
    result, raw_count = [], 0
    for values in iterator:
        raw_count += 1
        row = dict(zip(header, values))
        if row.get("Species (miRNA)") != "Homo sapiens" or row.get("Species (Target Gene)") != "Homo sapiens":
            continue
        left = mirna_symbol(str(row.get("miRNA", "")), candidates)
        right = resolve_symbol(str(row.get("Target Gene", "")), candidates, exact, aliases)
        if not left or not right:
            continue
        pmid = re.sub(r"\D", "", str(row.get("References (PMID)", "")))
        year = years.get(pmid)
        if year and year > 2022:
            continue
        result.append({
            "source": left, "target": right, "relation": "TARGETS",
            "database": "miRTarBase", "pmid": pmid, "year": year or "",
            "evidence_level": "strong_experimental",
            "mechanism": str(row.get("Experiments", "")), "direct": "",
            "release": "v9_2022",
        })
    return result, raw_count


def aggregate_and_cap(rows):
    """聚合同边多 PMID，并按直接性、证据数排序后限制每源最大边数。"""
    grouped = {}
    for row in rows:
        key = row["source"], row["target"], row["relation"], row["database"]
        item = grouped.setdefault(key, {**row, "pmids": set(), "years": []})
        if row["pmid"]:
            item["pmids"].add(row["pmid"])
        if row["year"]:
            item["years"].append(int(row["year"]))
        if row["direct"] == "t":
            item["direct"] = "t"
            item["evidence_level"] = "curated_causal_direct"
    by_source = defaultdict(list)
    for item in grouped.values():
        item["pmid"] = ";".join(sorted(item.pop("pmids"), key=lambda x: int(x)))
        year_values = item.pop("years")
        item["year"] = min(year_values) if year_values else ""
        item["pmid_count"] = len(item["pmid"].split(";")) if item["pmid"] else 0
        by_source[(item["source"], item["database"])].append(item)
    kept, dropped = [], 0
    for values in by_source.values():
        values.sort(key=lambda r: (
            r["direct"] != "t", -r["pmid_count"], r["target"], r["relation"]
        ))
        kept.extend(values[:MAX_PER_SOURCE_DB])
        dropped += max(0, len(values) - MAX_PER_SOURCE_DB)
    return kept, dropped, len(grouped)


def generate_features(edges, candidates):
    """由规范化边生成全局及逐数据库节点特征。

    逐源字段使模型可以分别做 SIGNOR/TRRUST/miRTarBase 消融；邻接到历史
    FerrDb promoter/suppressor 的计数则提供与预测任务直接相关的结构上下文。
    """
    roles = defaultdict(set)
    for edge in read_csv(GRAPH / "strict_historical_edges.csv"):
        if edge["source"] == "FerrDb":
            roles[edge["source_id"].split(":", 1)[1]].add(edge["relation"])
    values = defaultdict(Counter)
    neighbours = defaultdict(set)
    for edge in edges:
        source, target = edge["source"], edge["target"]
        database = edge["database"].lower()
        relation = edge["relation"].lower()
        neighbours[source].add(target)
        neighbours[target].add(source)
        for symbol, other, direction in ((source, target, "out"), (target, source, "in")):
            values[symbol][f"structured_{direction}_edge_count"] += 1
            values[symbol][f"structured_{edge['relation'].lower()}_{direction}_count"] += 1
            values[symbol][f"structured_{edge['database'].lower()}_edge_count"] += 1
            values[symbol]["structured_direct_count"] += int(edge["direct"] == "t")
            values[symbol]["structured_pmid_count"] += edge["pmid_count"]
            if "PROMOTES_FERROPTOSIS" in roles[other]:
                values[symbol]["structured_promoter_neighbour_count"] += 1
            if "SUPPRESSES_FERROPTOSIS" in roles[other]:
                values[symbol]["structured_suppressor_neighbour_count"] += 1
            values[symbol][f"{database}_{direction}_edge_count"] += 1
            values[symbol][f"{database}_{relation}_{direction}_count"] += 1
            values[symbol][f"{database}_pmid_count"] += edge["pmid_count"]
            values[symbol][f"{database}_direct_count"] += int(edge["direct"] == "t")
            if "PROMOTES_FERROPTOSIS" in roles[other]:
                values[symbol][f"{database}_promoter_neighbour_count"] += 1
            if "SUPPRESSES_FERROPTOSIS" in roles[other]:
                values[symbol][f"{database}_suppressor_neighbour_count"] += 1
    fields = [
        "gene", "structured_unique_neighbour_count", "structured_out_edge_count",
        "structured_in_edge_count", "structured_activates_out_count",
        "structured_inhibits_out_count", "structured_regulates_out_count",
        "structured_targets_out_count", "structured_signor_edge_count",
        "structured_trrust_edge_count", "structured_mirtarbase_edge_count",
        "structured_direct_count", "structured_pmid_count",
        "structured_promoter_neighbour_count",
        "structured_suppressor_neighbour_count",
    ]
    for database in ("signor", "trrust", "mirtarbase"):
        fields.extend([
            f"{database}_out_edge_count", f"{database}_in_edge_count",
            f"{database}_activates_out_count", f"{database}_inhibits_out_count",
            f"{database}_regulates_out_count", f"{database}_targets_out_count",
            f"{database}_pmid_count", f"{database}_direct_count",
            f"{database}_promoter_neighbour_count",
            f"{database}_suppressor_neighbour_count",
        ])
    rows = []
    for gene in sorted(candidates):
        row = {field: 0 for field in fields}
        row["gene"] = gene
        row["structured_unique_neighbour_count"] = len(neighbours[gene])
        for key, value in values[gene].items():
            if key in row:
                row[key] = value
        rows.append(row)
    return rows, fields


def main():
    """运行完整 ETL，并输出边、特征及审计报告。"""
    candidates, exact, aliases, uniprot = load_maps()
    years = pmid_years()
    signor, signor_raw = signor_edges(candidates, exact, aliases, uniprot, years)
    trrust, trrust_raw = trrust_edges(candidates, exact, aliases, years)
    mirtar, mirtar_raw = mirtarbase_edges(candidates, exact, aliases, years)
    edges, dropped, before_cap = aggregate_and_cap(signor + trrust + mirtar)
    edge_fields = [
        "source", "target", "relation", "database", "pmid", "pmid_count",
        "year", "evidence_level", "mechanism", "direct", "release",
    ]
    write_csv(GRAPH / "structured_historical_edges.csv", edges, edge_fields)
    features, feature_fields = generate_features(edges, candidates)
    write_csv(GRAPH / "structured_historical_features.csv", features, feature_fields)
    report = {
        "candidate_entities": len(candidates),
        "raw_rows": {
            "SIGNOR_Oct2022": signor_raw, "TRRUST_v2_2018": trrust_raw,
            "miRTarBase_v9_strong": mirtar_raw,
        },
        "mapped_rows_before_aggregation": {
            "SIGNOR": len(signor), "TRRUST": len(trrust), "miRTarBase": len(mirtar),
        },
        "unique_edges_before_cap": before_cap,
        "kept_edges": len(edges),
        "dropped_by_cap": dropped,
        "cap_per_source_database": MAX_PER_SOURCE_DB,
        "kept_by_database": dict(Counter(r["database"] for r in edges)),
        "kept_by_relation": dict(Counter(r["relation"] for r in edges)),
        "entities_with_structured_context": sum(
            bool(r["structured_unique_neighbour_count"]) for r in features
        ),
        "controls": [
            "Human only.",
            "Both endpoints must resolve into the bounded candidate universe.",
            "SIGNOR uses Oct2022 release; TRRUST uses 2018 v2; miRTarBase uses 2022 v9 strong evidence.",
            "PMID year > 2022 is excluded when year metadata is available.",
            "At most 50 edges per source entity per database after evidence-prioritized sorting.",
        ],
    }
    (GRAPH / "structured_history_report.json").write_text(
        json.dumps(report, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
